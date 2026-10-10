#!/usr/bin/env python3
"""Extract Bavaria 2025 all 96 Landkreis/kreisfreie Stadt completion figures.

Original Bavaria Landesamt F2200C 202500, table 4, four page spread sheets.
The original local region 3-digit keys are exact Bavaria Kreisschlüssel
(last three digits of the 5-digit AGS), prefix '09'. Join ONLY those 96 keys
against the 400 official BKG/HA26 districts; never join by ambiguous name.

The table distinguishes new residential building completed dwellings
(column 7) from completed dwellings under ALL construction measures
(column 23). Check seven Regierungsbezirk sums and published state total.
"""
from __future__ import annotations
import csv,json,re
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/"research/housing/raw/bavaria_f2200c_202500_official.xlsx"
OUT=ROOT/"research/housing/data";QA=ROOT/"research/housing/qa";WEB=ROOT/"topics/housing/data"
URL="https://www.statistik.bayern.de/mam/produkte/veroffentlichungen/statistische_berichte/f2200c_202500.xlsx"
def numeric(v,where):
 if v is None:return None
 if isinstance(v,(int,float)) and v>=0 and int(v)==v:return int(v)
 s=str(v).strip().replace(".","").replace(" ","")
 if s in ("-","/","x","","."):return None
 if not re.fullmatch(r"\d+",s):raise ValueError(f"Unrecognized official numeric value: {where}={v!r}")
 return int(s)
def main():
 atlas=json.loads((WEB/"atlas-counties.json").read_text(encoding="utf-8"))
 expected={x["id"] for x in atlas["counties"] if x["id"].startswith("09")}
 if len(expected)!=96:raise ValueError(f"Expected Bayern to comprise 96 Kreise got {len(expected)}")
 wb=load_workbook(SOURCE,read_only=True,data_only=True)
 rows={};region_sum={}
 for sh in wb:
  if not sh.title.startswith("Kreisübersicht Tab4"):continue
  for rn,r in enumerate(sh.iter_rows(values_only=True),start=1):
   if len(r)<23:continue
   key=str(r[0]).strip() if r[0] is not None else ""
   name=str(r[2]).strip() if r[2] is not None else ""
   if re.fullmatch(r"[1-7]",key) and name in (
     "Oberbayern","Niederbayern","Oberpfalz","Oberfranken",
     "Mittelfranken","Unterfranken","Schwaben"):
    if key in region_sum:raise ValueError(f"Repeated Bavarian regional summary {key}")
    region_sum[key]={"name":name,"new_dwellings":numeric(r[6],sh.title)}
   if not re.fullmatch(r"\d{3}",key):continue
   ags="09"+key
   if ags not in expected:raise ValueError(f"Unexpected Bavarian district code {ags}: {name}")
   if ags in rows:raise ValueError(f"Duplicate Bayern Kreis {ags} {sh.title}:{rn}")
   new_houses=numeric(r[4],(ags,"new houses"))
   new_units=numeric(r[6],(ags,"new residential dwellings"))
   new_small_houses=numeric(r[9],(ags,"new single/double houses"))
   new_small_units=numeric(r[11],(ags,"new single/double dwellings"))
   all_units=numeric(r[22],(ags,"all completed dwellings across measures"))
   if new_units is None or all_units is None or new_houses is None:
    raise ValueError(f"Missing official key count for {ags}")
   if new_units>all_units or new_small_units is not None and new_small_units>new_units:
    raise ValueError(f"Nested construction count impossible for {ags}: {new_units}/{all_units}")
   rows[ags]={"id":ags,"name":name,
      "completed_new_residential_buildings_2025":new_houses,
      "completed_dwellings_new_residential_buildings_2025":new_units,
      "completed_dwellings_new_small_residential_buildings_2025":new_small_units,
      "completed_dwellings_all_measures_2025":all_units}
 wb.close()
 if set(rows)!=expected:raise ValueError(f"2025 Bavaria official district coverage mismatch: missing={sorted(expected-set(rows))}, extra={sorted(set(rows)-expected)}")
 if len(region_sum)!=7:raise ValueError(f"Missing Bayern 7 Regierungsbezirk benchmarks: {region_sum}")
 region_check={}
 for k,v in sorted(region_sum.items()):
  data_sum=sum(x["completed_dwellings_new_residential_buildings_2025"] for key,x in rows.items() if key[2]==k)
  if data_sum!=v["new_dwellings"]:
   raise ValueError(f"REGION {k} {v['name']} county-sum {data_sum} != official Bezirks total {v['new_dwellings']}")
  region_check[v["name"]]=data_sum
 total_new=sum(x["completed_dwellings_new_residential_buildings_2025"] for x in rows.values())
 total_all=sum(x["completed_dwellings_all_measures_2025"] for x in rows.values())
 # The state's published headline includes residential/nonresidential and
 # changes to existing buildings. Confirm the all-measures counterpart.
 if total_all!=47359:raise ValueError(f"Official Bavaria 2025 published ALL completed dwellings 47,359 != {total_all}")
 if not 35_000<=total_new<=45_000:raise ValueError(f"New completed residential dwelling count implausible {total_new}")
 for d in (OUT,QA,WEB):d.mkdir(parents=True,exist_ok=True)
 payload={"meta":{"publisher":"Bayerisches Landesamt für Statistik",
    "reference_year":2025,"report":"F2200C 202500 Table 4",
    "original_url":URL,"counties":96,
    "new_residential_units":"Completed dwellings in NEW RESIDENTIAL buildings (column 7); directly comparable in definition to previously archived NRW new-residential category after considering local reporting methodology.",
    "all_finished_units":"ALL dwelling completions in new and existing buildings (column 23), including non-residential where dwellings are completed; NOT identical to new residential only.",
    "regional_control":"Sum of 96 districts checked exactly vs seven Regierungsbezirke, and vs official Bavaria 2025 overall total 47,359",
    "license":"Datenlizenz Deutschland Namensnennung 2.0",
    "provenance":"Official Bavarian 2025 annual report; 3-digit Bavarian district code prefixed by 09 exactly and validated vs 96 AGS."},
    "counties":[rows[k] for k in sorted(rows)]}
 (WEB/"bavaria-new-home-completions-2025.json").write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":"))+"\n",encoding="utf-8")
 with (OUT/"bavaria_2025_residential_completions_96_counties.csv").open("w",newline="",encoding="utf-8") as f:
  w=csv.DictWriter(f,fieldnames=list(rows[next(iter(rows))]));w.writeheader();w.writerows(payload["counties"])
 audit={"source":URL,"county_count":len(rows),"new_residential_finished_dwellings":total_new,
    "all_dwellings_finished":total_all,"bavaria_2025_official_all_dwellings":47359,
    "regierungsbezirke_checks":region_check,
    "notes":["Exact full administrative district code 09 + 3-digit Bayern region code",
     "Report 2025 is newer than 2023 nationwide regionalstatistik static archive",
     "Dwellings in new residential buildings and all completed dwellings are distinct measures"]}
 (QA/"bavaria_2025_96_district_completions_audit.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print("BAVARIA 2025 96 COUNTY PASS",total_new,"new-residential finished",total_all,"all completed dwellings",flush=True)
if __name__=="__main__":main()
