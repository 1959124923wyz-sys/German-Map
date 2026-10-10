#!/usr/bin/env python3
"""Archive official Saxony Zensus 2022 district, housing and household workbooks.

No guessed cell-to-county conversion. Separate inspection of sheet structure is
required before any extracted rates are published on a county map.
"""
from __future__ import annotations
import hashlib,json,time,urllib.request
from pathlib import Path
from openpyxl import load_workbook
BASE="https://www.zensus.sachsen.de/download/"
FILES={
 "saxony_zensus2022_county_buildings.xlsx":"04_Rueckblick%20Zensus%202011/statistik-sachsen_regionaltabelle_Z22_gebaeude-wohnungen_kreis-sn.xlsx",
 "saxony_zensus2022_county_households.xlsx":"04_Rueckblick%20Zensus%202011/statistik-sachsen_regionaltabelle_Z22_haushalte_kreis-sn.xlsx",
 "saxony_zensus2022_all_buildings.xlsx":"04_Rueckblick%20Zensus%202011/statistik-sachsen_zensus_2022_FI_GWZ_1.xlsx",
 "saxony_zensus2022_vacancy.xlsx":"04_Gebaeude%20und%20Wohnungszaehlung/statistik-sachsen_zensus_2022_FI_GWZ_2.xlsx",
}
ROOT=Path(__file__).resolve().parents[2]
RAW=ROOT/"research/housing/raw";QA=ROOT/"research/housing/qa"
SOURCE="https://www.zensus.sachsen.de/zensus-2022.html"
def download(url):
 for i in range(3):
  try:
   req=urllib.request.Request(url,headers={"User-Agent":"Mozilla/5.0 (compatible; GermanMap/1.0; public statistical archival research)"})
   with urllib.request.urlopen(req,timeout=120) as resp:
    b=resp.read()
   if len(b)<5000 or not b.startswith(b"PK"):raise ValueError(f"Invalid official XLSX ({len(b)} bytes)")
   return b
  except Exception as e:
   print("RETRY",i,url,type(e).__name__,str(e),flush=True)
   if i==2:raise
   time.sleep(3*(i+1))
def schema(file):
 wb=load_workbook(file,read_only=True,data_only=True)
 sheets={}
 for sh in wb:
  sample=[]
  for r in sh.iter_rows(max_row=min(sh.max_row,22),max_col=min(sh.max_column,18),values_only=True):
   sample.append([str(v)[:150] if v is not None else None for v in r])
  sheets[sh.title]={"rows":sh.max_row,"columns":sh.max_column,"sample_first_22_rows":sample}
 wb.close()
 return sheets
def main():
 RAW.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
 audit={"state":"Sachsen","publisher":"Statistisches Landesamt des Freistaates Sachsen",
   "landing":SOURCE,"reference_date":"2022-05-15","purpose":"Raw source archival; never impute suppressed cells; figures use 2022 administrative boundaries",
   "files":{}}
 for name,path in FILES.items():
  url=BASE+path;b=download(url);file=RAW/name;file.write_bytes(b)
  a={"url":url,"bytes":len(b),"sha256":hashlib.sha256(b).hexdigest(),"sheets":schema(file)}
  audit["files"][name]=a
  print("VERIFIED",name,len(b),"bytes",len(a["sheets"]),"sheets",flush=True)
 (QA/"saxony_zensus2022_workbook_inventory.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print("SUCCESS",len(FILES),"official Saxony Zensus 2022 original workbooks stored.",flush=True)
if __name__=="__main__":main()
