#!/usr/bin/env python3
"""Archive Brandenburg's original 2025 Kreis Baufertigstellungen official XLSX.

2026 press release includes 'Gemeldete fertiggestellte Wohnungen 2025 nach
Kreisen und kreisfreien Städten' linked original publication XLSX.
Stage1 archives and records exact header columns/rows, NOT a guessed extract.
"""
from __future__ import annotations
import hashlib,json,urllib.request
from pathlib import Path
from openpyxl import load_workbook
URL="https://download.statistik-berlin-brandenburg.de/94c8873070894dca/43d5f940c81b/pressemitteilung-tabelle-60-2026.xlsx"
LANDING="https://www.statistik-berlin-brandenburg.de/presse/2026/60-baufertigstellungen-2025-brandenburg/"
ROOT=Path(__file__).resolve().parents[2]
RAW=ROOT/"research/housing/raw";QA=ROOT/"research/housing/qa"
def main():
 RAW.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
 req=urllib.request.Request(URL,headers={"User-Agent":"GermanMap open statistical data research"})
 with urllib.request.urlopen(req,timeout=85) as resp:
  raw=resp.read(18_000_000)
  print("BBAW",resp.status,len(raw),resp.headers.get("Content-Type"),flush=True)
 if len(raw)<5000 or not raw.startswith(b"PK") or len(raw)>=18_000_000:raise ValueError("Invalid Brandenburg official XLSX")
 source=RAW/"brandenburg_2025_18_county_completions_official.xlsx";source.write_bytes(raw)
 w=load_workbook(source,read_only=True,data_only=True)
 report={"publisher":"Amt für Statistik Berlin-Brandenburg","statistical_year":2025,"land":"Brandenburg",
         "download":URL,"landing":LANDING,"original_bytes":len(raw),
         "sha256":hashlib.sha256(raw).hexdigest(),"sheets":{},
         "reference_press_state_total_completed_dwellings_all_measures":7379,
         "reference_press_state_total_new_dwellings_in_residential_buildings":6457,
         "note":"Two different completion category totals (all/new residential). Do not confuse 7379 with 6457."}
 for sh in w:
  lines=[]
  for row in sh.iter_rows(max_row=min(40,sh.max_row),max_col=min(28,sh.max_column),values_only=True):
   lines.append([str(x)[:145] if x is not None else None for x in row])
  report["sheets"][sh.title]={"rows":sh.max_row,"cols":sh.max_column,"first40":lines}
  print("SHEET",sh.title,sh.max_row,sh.max_column,flush=True)
 w.close()
 (QA/"brandenburg_2025_completions_original_schema.json").write_text(
  json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print("ARCHIVED official Brandenburg 2025 county construction XLSX",flush=True)
if __name__=="__main__":main()
