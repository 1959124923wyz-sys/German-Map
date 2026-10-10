#!/usr/bin/env python3
"""2025 Brandenburg completed dwellings: exact official press XLSX, all 18 Kreise.

Official workbook has no printed territorial codes, unlike raw Bavaria. To
avoid fuzzy-name joins, use an explicit 18-entry 2024 AGS crosswalk and REQUIRE
exact identity of the names against the current official HA26 county table.
Any missing/unexpected/duplicate names abort output.

IMPORTANT different construction scope: 'Neubau' in the official 2025 press
table sums new residential + new non-residential buildings (6489),
NOT new RESIDENTIAL buildings only (6457). So data is published as separate
Brandenburg-only county detail, not combined into a homogeneous 2025 choropleth.
"""
from __future__ import annotations
import csv,json,re
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/"research/housing/raw/brandenburg_2025_18_county_completions_official.xlsx"
WEB=ROOT/"topics/housing/data";DATA=ROOT/"research/housing/data";QA=ROOT/"research/housing/qa"
URL="https://download.statistik-berlin-brandenburg.de/94c8873070894dca/43d5f940c81b/pressemitteilung-tabelle-60-2026.xlsx"
LANDING="https://www.statistik-berlin-brandenburg.de/presse/2026/60-baufertigstellungen-2025-brandenburg/"
# Explicit unique source names in original AfS table and canonical 2024 county AGS.
CROSSWALK={
 "Brandenburg an der Havel, Stadt":"12051",
 "Cottbus, Stadt":"12052","Frankfurt (Oder), Stadt":"12053","Potsdam, Stadt":"12054",
 "Barnim":"12060","Dahme-Spreewald":"12061","Elbe-Elster":"12062",
 "Havelland":"12063","Märkisch-Oderland":"12064","Oberhavel":"12065",
 "Oberspreewald-Lausitz":"12066","Oder-Spree":"12067",
 "Ostprignitz-Ruppin":"12068","Potsdam-Mittelmark":"12069",
 "Prignitz":"12070","Spree-Neiße":"12071","Teltow-Fläming":"12072","Uckermark":"12073"}
def number(v):
 if isinstance(v,int) and v>=0:return v
 s=str(v).strip()
 if re.fullmatch(r"\d+",s):return int(s)
 raise ValueError(f"Expected full nonnegative count; got {v!r}")
def main():
 atlas=json.loads((WEB/"atlas-counties.json").read_text(encoding="utf-8"))
 official={r["id"]:r["name"].strip() for r in atlas["counties"] if r["id"].startswith("12")}
 if len(official)!=18 or set(CROSSWALK.values())!=set(official):raise ValueError("Canonical 18 official AGS mismatch")
 for name,key in CROSSWALK.items():
  if name!=official[key]:raise ValueError(f"Official HA26 name changed, do not silently name-join: {name} != {official[key]}")
 book=load_workbook(SRC,read_only=True,data_only=True)
 sh=book["Tabelle1"]
 rows=list(sh.iter_rows(values_only=True))
 if "Neubau" not in str(rows[4][3]) or rows[5][1]!=2025 and str(rows[5][1])!="2025":
  raise ValueError("Brandenburg printed year/category headers changed")
 state=rows[7]
 if str(state[0]).strip()!="Brandenburg":raise ValueError("Missing 2025 Brandenburg official total row")
 totals=[number(state[j]) for j in (1,2,3,4)]
 if totals!=[7379,10172,6489,9466]:raise ValueError(f"Unexpected official 2025/24 totals {totals}")
 records={}
 for row in rows[8:]:
  name=str(row[0]).strip() if row[0] is not None else ""
  if name not in CROSSWALK:
   if not name or name.startswith(("1 ","Quelle:","Diese Seite steht")):continue
   raise ValueError(f"Unexpected non-county data row {name!r}")
  key=CROSSWALK[name]
  if key in records:raise ValueError(f"Duplicate official Brandenburg Kreis {key}")
  cells=[number(row[i]) for i in (1,2,3,4)]
  # NOTE: New-build count can exceed total if demolition is netted out for
  # all measures; do not incorrectly assert all >= new (source footnote).
  records[key]={"id":key,"name":name,
   "completed_dwellings_all_measures_2025":cells[0],
   "completed_dwellings_all_measures_2024":cells[1],
   "completed_new_build_all_types_2025":cells[2],
   "completed_new_build_all_types_2024":cells[3]}
 if set(records)!=set(official):raise ValueError(f"Not all 18 Brandenburg Kreise parsed: {sorted(set(official)-set(records))}")
 actual=[sum(r[field] for r in records.values()) for field in
        ("completed_dwellings_all_measures_2025","completed_dwellings_all_measures_2024",
         "completed_new_build_all_types_2025","completed_new_build_all_types_2024")]
 if actual!=totals:raise ValueError(f"Brandenburg county sums fail published 4 controls: {actual} vs {totals}")
 for directory in (WEB,DATA,QA):directory.mkdir(parents=True,exist_ok=True)
 payload={"meta":{"publisher":"Amt für Statistik Berlin-Brandenburg","source":URL,
   "press_page":LANDING,"region":"Land Brandenburg","districts":18,
   "years":[2024,2025],"all_completions_2025_state_total":7379,
   "new_build_all_types_2025_state_total":6489,
   "new_residential_only_state_total_2025_from_press":6457,
   "warning":"Neubau 6,489 includes new RESIDENTIAL and new NON-RESIDENTIAL buildings, NOT comparable in coverage with Bavaria / NRW new-residential-only counts. All-measures 7,379 counts all completed housing incl existing measures.",
   "name_key_method":"Only all-18 exact unique AfS names are matched to separately verified official HA26 2024 5-digit Kreis AGS, otherwise fail.",
   "license":"CC BY 3.0 Germany, Amt für Statistik Berlin-Brandenburg"},
  "counties":[records[key] for key in sorted(records)]}
 (WEB/"brandenburg-completions-2024-2025.json").write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":"))+"\n",encoding="utf-8")
 with (DATA/"brandenburg_2024_2025_18_county_completions.csv").open("w",encoding="utf-8",newline="") as f:
  w=csv.DictWriter(f,fieldnames=list(records[next(iter(records))]));w.writeheader();w.writerows(payload["counties"])
 (QA/"brandenburg_2025_18_county_completions_audit.json").write_text(json.dumps({
    "source":URL,"report_footnote":"New-build completion count may exceed total due net housing change from other building measures",
    "counts":len(records),"source_state_totals":totals,"county_sums":actual,
    "crosswalk":CROSSWALK,"license":"CC BY 3.0 Germany",
    "warning":"This report has 6489 new-build homes including nonresidential; its new residential-only total 6457 cannot be assigned across counties from this press table without additional source"
  },ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print("BRANDENBURG official 2024/2025 18 county PASS",actual,flush=True)
if __name__=="__main__":main()
