#!/usr/bin/env python3
"""Acquire raw GREIX 2026 Q2 city workbook and Zensus 2022 local existing rents.

This is a research acquisition/inspection pass. No city-to-AGS inference.
Both providers require source attribution. Run with openpyxl installed.
"""
from __future__ import annotations
import csv, hashlib, io, json, re, time, urllib.request
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
RAW=ROOT/"research/housing/raw"
QA=ROOT/"research/housing/qa"
DATA=ROOT/"research/housing/data"
GREIX="https://www.kielinstitut.de/fileadmin/Dateiverwaltung/IfW_Unit/Macroeconomics/GREIX/Mietpreisindex/City_metrics_public.xlsx"
EXISTING="https://deutschlandatlas.bund.de/fileadmin/Downloaddateien/HA26/Deutschlandatlas_VBGEM1222_HA26.csv"

def download(url):
    for n in range(3):
        try:
            q=urllib.request.Request(url,headers={"User-Agent":"Mozilla/5.0 (GermanMap research, source attribution preserved)"})
            with urllib.request.urlopen(q,timeout=80) as f:
                raw=f.read()
                print("DOWNLOADED",f.status,len(raw),url)
                if len(raw)<10000: raise ValueError("Suspiciously small response")
                return raw
        except Exception as e:
            print("RETRY",n,type(e).__name__,str(e))
            if n==2:raise
            time.sleep(3+4*n)

def workbook_report(path):
    book=load_workbook(path,read_only=True,data_only=True)
    report={"source":GREIX,"source_credit":"Kiel Institut für Weltwirtschaft auf Basis der VALUE Marktdatenbank",
            "release":"Q2/2026, published 2026-07-22","sheets":{}}
    for sheet in book:
        head=list(sheet.iter_rows(min_row=1,max_row=min(9,sheet.max_row),
                       max_col=min(12,sheet.max_column),values_only=True))
        report["sheets"][sheet.title]={"rows":sheet.max_row,"columns":sheet.max_column,
          "first_rows":[[str(v)[:100] if v is not None else None for v in row] for row in head]}
        print("SHEET",sheet.title, sheet.max_row, sheet.max_column,"HEAD",report["sheets"][sheet.title]["first_rows"][:2])
    book.close()
    return report

def municipal_rents(raw):
    try:
        txt=raw.decode("utf-8-sig")
    except UnicodeError:
        txt=raw.decode("cp1252")
    if "preis_miet_best" not in txt:
        raise ValueError("Missing 2022 rent indicator in decoded CSV")
    rows=csv.reader(io.StringIO(txt),delimiter=";")
    hdr=[h.strip().lstrip("\ufeff").lower() for h in next(rows)]
    print("VBGEM HEAD",hdr)
    try: idx=hdr.index("preis_miet_best")
    except ValueError: raise ValueError("Existing rent field not in official municipal-association file")
    records=[]; discarded=[]
    for line,row in enumerate(rows,start=2):
        if len(row)<=idx:continue
        key=row[0].strip()
        if not re.fullmatch(r"\d{5,12}",key):
            if len(discarded)<5:discarded.append([line,*row[:3]])
            continue
        value=row[idx].strip().replace(",",".")
        if value in ("","-9999","-99999","x",".","/","-"): number=None
        else:
            try: number=float(value)
            except ValueError:number=None
        if number is not None and not 1<=number<=65:
            raise ValueError(f"Suspicious 2022 monthly rent per sq m: {key} {number}")
        records.append({"regional_key_raw":key,"name":row[1].strip(),"net_cold_rent_2022_eur_m2":number})
    if len(records)<500:raise ValueError(f"Unexpected count {len(records)}")
    return records,{"source":EXISTING,"field":"preis_miet_best","unit":"EUR/m2","year":"2022",
     "level":"Gemeindeverband: NOT uniform county-level data",
     "count":len(records),"with_rent":sum(x["net_cold_rent_2022_eur_m2"] is not None for x in records),
     "key_lengths":{str(n):sum(len(x["regional_key_raw"])==n for x in records)
                    for n in sorted({len(x["regional_key_raw"]) for x in records})},
     "sample":records[:8],"discarded":discarded}

def main():
    for d in (RAW,QA,DATA):d.mkdir(parents=True,exist_ok=True)
    files=[
      ("greix_city_metrics_q2_2026.xlsx",GREIX),
      ("deutschlandatlas_vbgem1222_ha26.csv",EXISTING)
    ]
    manifests={}
    for name,url in files:
        raw=download(url);p=RAW/name;p.write_bytes(raw)
        manifests[name]={"url":url,"sha256":hashlib.sha256(raw).hexdigest(),"bytes":len(raw)}
    (QA/"greix_workbook_structure.json").write_text(json.dumps(
      workbook_report(RAW/files[0][0]),ensure_ascii=False,indent=2,default=str)+"\n",encoding="utf-8")
    records,meta=municipal_rents((RAW/files[1][0]).read_bytes())
    with (DATA/"vbgem_existing_rents_2022.csv").open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=list(records[0]));w.writeheader();w.writerows(records)
    (QA/"vbgem_rent_audit.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    (QA/"secondary_raw_manifest.json").write_text(json.dumps(manifests,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("DONE",len(records),"municipal association rows and GREIX raw workbook.")

if __name__=="__main__":main()
