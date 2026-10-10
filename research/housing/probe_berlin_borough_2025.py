#!/usr/bin/env python3
"""Acquire 2025 Berlin Borough reported housing completions official revised XLSX.

Important: 2026-05-22 correction to Berlin press release 59/2026. Use ONLY
corrected workbook, not the withdrawn original or general Berlin sums.
Boroughs are a subcounty geography: never treat borough data as 12 Landkreis.
"""
from __future__ import annotations
import hashlib,json,urllib.request
from pathlib import Path
from openpyxl import load_workbook
URL="https://download.statistik-berlin-brandenburg.de/f04f3ade7f98cd05/090fd40a76ac/pressemitteilung-tabelle-59_2026.xlsx"
LANDING="https://www.statistik-berlin-brandenburg.de/presse/2026/59-baufertigstellungen-2025-berlin/"
ROOT=Path(__file__).resolve().parents[2]
RAW=ROOT/"research/housing/raw";QA=ROOT/"research/housing/qa"
def main():
 RAW.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
 req=urllib.request.Request(URL,headers={"User-Agent":"GermanMap public official data archival study"})
 with urllib.request.urlopen(req,timeout=85) as response:
  raw=response.read(18_000_000)
  print("FETCH",response.status,len(raw),flush=True)
 if not raw.startswith(b"PK") or not (5000<len(raw)<18_000_000):raise ValueError("Not credible corrected official Excel")
 path=RAW/"berlin_2025_borough_completions_corrected_official.xlsx";path.write_bytes(raw)
 wb=load_workbook(path,read_only=True,data_only=True)
 details={}
 for sh in wb:
  preview=[]
  for row in sh.iter_rows(max_row=min(sh.max_row,48),max_col=min(sh.max_column,25),values_only=True):
   preview.append([str(x)[:120] if x is not None else None for x in row])
  details[sh.title]={"rows":sh.max_row,"cols":sh.max_column,"first48":preview}
 wb.close()
 audit={"source":URL,"landing":LANDING,"source_correction":"Corrected official Berlin 2025 59/2026 press release, May 22 2026",
        "territorial_level":"12 Berlin Bezirke (NOT Kreise); Berlin as whole is one Kreis 11000",
        "official_total_all_completed_dwellings_2025":11027,
        "official_total_new_building_completed_dwellings_2025":9524,
        "bytes":len(raw),"sha256":hashlib.sha256(raw).hexdigest(),
        "sheets":details}
 (QA/"berlin_2025_corrected_borough_completions_schema.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print("ARCHIVED BERLIN 2025 CORRECTED XLSX",list(details),flush=True)
if __name__=="__main__":main()
