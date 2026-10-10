#!/usr/bin/env python3
"""Berlin's 2024/2025 completions by all 12 boroughs, corrected original XLSX.

Berlin's 12 Bezirke are internal administrative districts, NOT 12 separate
Landkreise; the national map still contains exactly ONE Berlin Kreis 11000.
The officially corrected 59/2026 XLSX must be used; confirm all four borough
aggregates against published Berlin-wide benchmarks.
"""
from __future__ import annotations
import csv,json,re
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
FILE=ROOT/"research/housing/raw/berlin_2025_borough_completions_corrected_official.xlsx"
WEB=ROOT/"topics/housing/data";DATA=ROOT/"research/housing/data";QA=ROOT/"research/housing/qa"
URL="https://download.statistik-berlin-brandenburg.de/f04f3ade7f98cd05/090fd40a76ac/pressemitteilung-tabelle-59_2026.xlsx"
PAGE="https://www.statistik-berlin-brandenburg.de/presse/2026/59-baufertigstellungen-2025-berlin/"
BEZIRKE=[
 "Mitte","Friedrichshain-Kreuzberg","Pankow","Charlottenburg-Wilmersdorf",
 "Spandau","Steglitz-Zehlendorf","Tempelhof-Schöneberg","Neukölln",
 "Treptow-Köpenick","Marzahn-Hellersdorf","Lichtenberg","Reinickendorf"]
def asint(v):
 s=str(v).strip()
 if not re.fullmatch(r"\d+",s):raise ValueError(f"Invalid corrected Berlin dwelling count {v!r}")
 return int(s)
def main():
 wb=load_workbook(FILE,read_only=True,data_only=True)
 sh=wb["Tabelle1"];rows=list(sh.iter_rows(values_only=True));wb.close()
 if str(rows[5][1])!="2025" or "Neubau" not in str(rows[4][3]):raise ValueError("Berlin corrected source columns changed")
 if str(rows[7][0]).strip()!="Berlin":raise ValueError("Missing Berlin entire city row")
 controls=[asint(rows[7][i]) for i in range(1,5)]
 if controls!=[11027,15362,9524,14632]:raise ValueError(f"Not corrected Berlin data: {controls}")
 out=[]
 for i,name in enumerate(BEZIRKE,start=10):
  r=rows[i-1]
  if str(r[0]).strip()!=name:raise ValueError(f"Berlin Borough name/order mismatch at row {i}: {r[0]!r} vs {name}")
  a=[asint(r[j]) for j in range(1,5)]
  out.append({"borough_index":i-9,"name":name,
   "completed_dwellings_all_measures_2025":a[0],
   "completed_dwellings_all_measures_2024":a[1],
   "completed_dwellings_new_build_all_types_2025":a[2],
   "completed_dwellings_new_build_all_types_2024":a[3]})
 if len(out)!=12:raise ValueError("Not 12 Berlin Bezirke")
 sums=[sum(row[k] for row in out) for k in (
   "completed_dwellings_all_measures_2025","completed_dwellings_all_measures_2024",
   "completed_dwellings_new_build_all_types_2025","completed_dwellings_new_build_all_types_2024")]
 if sums!=controls:raise ValueError(f"12 borough sums {sums} != 4 corrected official totals {controls}")
 for p in (WEB,DATA,QA):p.mkdir(parents=True,exist_ok=True)
 data={"meta":{"publisher":"Amt für Statistik Berlin-Brandenburg",
  "official_download":URL,"press_page":PAGE,"press_release":"59/2026, correction May 22 2026",
  "geography":"12 Berlin Bezirke (SUBCOUNTY units), NOT Kreise","berlin_one_county_id":"11000",
  "reference_years":[2024,2025],"borough_count":12,
  "all_completed_dwellings_2025":11027,
  "new_build_all_type_completed_dwellings_2025":9524,
  "warning":"New-build category includes dwellings in nonresidential buildings and must not be equated to NRW/Bavaria new-residential-only subset",
  "license":"CC BY 3.0 DE (Amt für Statistik Berlin-Brandenburg)"},
  "boroughs":out}
 (WEB/"berlin-2024-2025-borough-completions.json").write_text(json.dumps(data,ensure_ascii=False,separators=(",",":"))+"\n",encoding="utf-8")
 with (DATA/"berlin_2024_2025_12_borough_completions.csv").open("w",encoding="utf-8",newline="") as f:
  w=csv.DictWriter(f,fieldnames=list(out[0]));w.writeheader();w.writerows(out)
 (QA/"berlin_2025_corrected_12_borough_completions_audit.json").write_text(json.dumps({
    "source":URL,"source_is_corrected":True,
    "borough_count":len(out),"official_four_berlin_totals":controls,
    "borough_four_totals":sums,
    "note":"Berlin Bezirke are subcounty; do not generate false extra county polygons/AGS or combine new-build with all-measures."},ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print("BERLIN 2025 REVISED OFFICIAL",len(out),"boroughs four controls MATCH",sums,flush=True)
if __name__=="__main__":main()
