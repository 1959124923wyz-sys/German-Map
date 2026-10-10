#!/usr/bin/env python3
"""Archive Germany's nationwide 2022 housing census regional XLSX and inventory.

Primary original source: Destatis Zensus 2022, 20 MB full Regionaltabelle
Gebäude und Wohnungen, including municipality / administrative association /
Landkreis levels. No claims about county metrics before inspecting workbook.
Avoid pushing an entire workbook dump into QA; retain sha256 and concise schema.
"""
from __future__ import annotations
import hashlib,json,time,urllib.request
from pathlib import Path
from openpyxl import load_workbook
URL="https://www.destatis.de/static/DE/zensus/gitterdaten/Regionaltabelle_Gebaeude_Wohnungen.xlsx"
LANDING="https://www.destatis.de/DE/Themen/Gesellschaft-Umwelt/Bevoelkerung/Zensus2022/_inhalt.html"
ROOT=Path(__file__).resolve().parents[2]
RAW=ROOT/"research/housing/raw";QA=ROOT/"research/housing/qa"
def fetch():
 for i in range(3):
  try:
   req=urllib.request.Request(URL,headers={"User-Agent":"Mozilla/5.0 GermanMap census archival research","Accept":"*/*"})
   with urllib.request.urlopen(req,timeout=160) as f:
    data=f.read(60_000_000)
    print("FETCH",f.status,len(data),"bytes",f.headers.get("Content-Type"),flush=True)
   if not (10_000_000<len(data)<60_000_000) or not data.startswith(b"PK"):
    raise ValueError("Unexpected official 20 MB national census workbook")
   return data
  except Exception as e:
   print("RETRY",i,type(e).__name__,str(e)[:200],flush=True)
   if i==2:raise
   time.sleep(5*(i+1))
def inventory(path):
 book=load_workbook(path,read_only=True,data_only=True)
 sheets={}
 for sh in book:
  nonempty=[]
  for index,row in enumerate(sh.iter_rows(min_row=1,max_row=min(sh.max_row,20),
                  max_col=min(sh.max_column,16),values_only=True),start=1):
   nonempty.append({"row":index,"cells":[str(v)[:110] if v is not None else "" for v in row]})
  sheets[sh.title]={"rows":sh.max_row,"cols":sh.max_column,"first_20_rows":nonempty}
  print("SHEET",sh.title,sh.max_row,sh.max_column,flush=True)
 book.close()
 return sheets
def main():
 RAW.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
 raw=fetch();name="destatis_zensus_2022_regional_housing_national.xlsx"
 path=RAW/name;path.write_bytes(raw)
 item={"url":URL,"landing":LANDING,"publisher":"Statistische Ämter des Bundes und der Länder",
       "territorial_date":"2022-05-15","license":"Data Licence Germany Attribution 2.0",
       "bytes":len(raw),"sha256":hashlib.sha256(raw).hexdigest(),
       "note":"No county extraction yet. Some rows aggregate municipalities/associations/counties and MUST NOT be summed as if independent.",
       "sheets":inventory(path)}
 (QA/"destatis_zensus2022_national_housing_workbook_schema.json").write_text(
  json.dumps(item,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 print("ARCHIVED NATIONAL 2022 ZENSUS HOUSING",len(raw),"bytes",flush=True)
if __name__=="__main__":main()
