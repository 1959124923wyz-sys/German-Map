#!/usr/bin/env python3
"""Document exact human-readable labels for the machine column codes in the
national official Zensus 2022 regional housing table. No inferred semantics.
"""
from __future__ import annotations
import json
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/"research/housing/raw/destatis_zensus_2022_regional_housing_national.xlsx"
OUT=ROOT/"research/housing/qa/zensus2022_official_housing_column_codebook.json"
def main():
 wb=load_workbook(SRC,read_only=True,data_only=True)
 out={"source":"https://www.destatis.de/static/DE/zensus/gitterdaten/Regionaltabelle_Gebaeude_Wohnungen.xlsx",
      "year":2022,"rows":{}}
 for sheet,raw in [("dwellings","Wohnungen"),("buildings","Gebäude ")]:
  csvsheet=wb["CSV-Wohnungen" if sheet=="dwellings" else "CSV-Gebäude"]
  h=[str(x or "") for x in next(csvsheet.iter_rows(max_row=1,values_only=True))]
  shown=wb[raw]
  sample=list(shown.iter_rows(min_row=3,max_row=10,values_only=True))
  codes={}
  for i,code in enumerate(h[4:]):
   human_col=i+3
   path=[]
   for r in sample:
    val=r[human_col] if human_col<len(r) else None
    if val is not None and str(val).strip():path.append(str(val).strip())
   codes[code]={"column":i+4,"labels":path}
  out["rows"][sheet]=codes
 wb.close()
 OUT.write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
 for s in out["rows"]:
  print("SHEET",s,"FIELDS",len(out["rows"][s]),flush=True)
  for code,labels in out["rows"][s].items():
   if any(k in code for k in ["LEERSTAND","QMMIETE","MIETE_EURM","LEQ","ETQ","NUTZUNG"]):
    print(code,labels["labels"],flush=True)
if __name__=="__main__":main()
