#!/usr/bin/env python3
"""Reproducibly transform GREIX public monthly city price workbook into typed series.

Uses the July 2026 Q2 workbook previously archived in raw/. It deliberately
keeps nominal data only; inflation-adjusted and nominal observations must not
be conflated. City names are not mapped to county AGS by guessing.
"""
from __future__ import annotations
import csv, json
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
INPUT=ROOT/"research/housing/raw/greix_city_metrics_q2_2026.xlsx"
OUT=ROOT/"research/housing/data"
WEB=ROOT/"topics/housing/data"
QA=ROOT/"research/housing/qa"
URL="https://www.kielinstitut.de/fileadmin/Dateiverwaltung/IfW_Unit/Macroeconomics/GREIX/Mietpreisindex/City_metrics_public.xlsx"
COLUMNS=("city","year","quarter","month","rent_avg_eur_m2","rent_median_eur_m2","rent_p25_eur_m2","rent_p75_eur_m2","hedonic_index")

def positive(v):
    if v is None or v=="":return None
    n=float(v)
    if n<=0:return None
    return round(n,4)

def main():
    if not INPUT.exists(): raise FileNotFoundError(f"Run secondary intake first: {INPUT}")
    book=load_workbook(INPUT,read_only=True,data_only=True)
    sheet=book["Tabelle1"]
    rows=sheet.iter_rows(values_only=True)
    head=[str(s).strip() for s in next(rows)]
    index={h:i for i,h in enumerate(head)}
    must=("Year","Quarter","Month","City","Index","AVG_PRICE_SQM","MED_PRICE_SQM",
          "P75_PRICE_SQM","P25_PRICE_SQM","Inflation_adjusted")
    if set(must)-set(index):
        raise ValueError(f"Missing GREIX columns: {set(must)-set(index)}")
    result=[]; seen=set(); skips=Counter()
    for rn,raw in enumerate(rows,start=2):
        def get(s):return raw[index[s]]
        is_real=get("Inflation_adjusted")
        if str(is_real).strip() not in ("0","0.0","False"):
            skips["inflation_adjusted"]+=1
            continue
        if get("City") is None:
            skips["no_city"]+=1
            continue
        city=str(get("City")).strip()
        year,month,quarter=map(int,(get("Year"),get("Month"),get("Quarter")))
        if not (2012<=year<=2026 and 1<=month<=12 and 1<=quarter<=4):
            raise ValueError(f"Invalid time {year}-{month}/{quarter}")
        if (month-1)//3+1!=quarter:
            raise ValueError(f"Quarter/month mismatch {year}-{month}/{quarter}")
        key=(city,year,month)
        if key in seen: raise ValueError(f"Duplicated nominal monthly observation {key}")
        seen.add(key)
        d=dict(zip(COLUMNS,(city,year,quarter,month,
             positive(get("AVG_PRICE_SQM")),
             positive(get("MED_PRICE_SQM")),
             positive(get("P25_PRICE_SQM")),
             positive(get("P75_PRICE_SQM")),
             positive(get("Index")))))
        if d["rent_avg_eur_m2"] is None or not 1<d["rent_avg_eur_m2"]<120:
            raise ValueError(f"Missing/extreme GREIX nominal price: {rn} {d}")
        if d["rent_p25_eur_m2"] is not None and d["rent_p75_eur_m2"] is not None and d["rent_p25_eur_m2"]>d["rent_p75_eur_m2"]:
            raise ValueError(f"Inverted quartiles {rn} {d}")
        result.append(d)
    book.close()
    result.sort(key=lambda d:(d["city"].casefold(),d["year"],d["month"]))
    cities=defaultdict(list)
    for d in result:
        cities[d["city"]].append([f'{d["year"]:04d}-{d["month"]:02d}',d["rent_avg_eur_m2"],
                                  d["hedonic_index"],d["rent_median_eur_m2"],
                                  d["rent_p25_eur_m2"],d["rent_p75_eur_m2"]])
    if not 25<=len(cities)<=120 or len(result)<5000:
        raise ValueError(f"Implausible GREIX workbook coverage cities={len(cities)} rows={len(result)}")
    QA.mkdir(parents=True,exist_ok=True);WEB.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
    with (OUT/"greix_nominal_monthly_city_series.csv").open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=COLUMNS);w.writeheader();w.writerows(result)
    full={"meta":{"source":URL,"credit":"Kiel Institut für Weltwirtschaft auf Basis der VALUE Marktdatenbank",
                  "release":"2026-07-22 Q2/2026",
                  "type":"nominal GREIX city/region asking rents; not official nationwide county data",
                  "time_frequency":"monthly",
                  "columns":["YYYY-MM","avg EUR/m2","hedonic_index","median EUR/m2","p25 EUR/m2","p75 EUR/m2"],
                  "city_count":len(cities),"observations":len(result)},
          "cities":dict(sorted(cities.items()))}
    (WEB/"greix-city-series.json").write_text(json.dumps(full,ensure_ascii=False,separators=(",",":"))+"\n",encoding="utf-8")
    audit={"source":URL,"cities":len(cities),"observations_nominal":len(result),
           "skipped":dict(skips),"latest":max(f'{d["year"]:04d}-{d["month"]:02d}' for d in result),
           "by_city":{city:{"first":r[0][0],"last":r[-1][0],"months":len(r)} for city,r in sorted(cities.items())},
           "important":"2025 annual BBSR rents and Q2 2026 GREIX monthly city data are different sampling/methodological series and should never be described as the same time series."}
    (QA/"greix_nominal_audit.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("GREIX OK",len(cities),"cities/regions",len(result),"monthly nominal observations",audit["latest"])
    print("CITIES",list(cities))

if __name__=="__main__":
    main()
