#!/usr/bin/env python3
"""Parse Saxony's official housed-homeless count for its 13 counties, 2022-26.

The source round-counts to the nearest 5 for disclosure control, and totals
need not be exactly additive. This is NOT all homelessness or all of Germany.
Explicit county AGS crosswalk audited against the old existing map.
"""
from __future__ import annotations
import csv,json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
RAW=ROOT/"research/housing/raw"
OUT=ROOT/"research/housing/data"
QA=ROOT/"research/housing/qa"
WEB=ROOT/"topics/housing/data"
NAMES={
 "Chemnitz, Stadt":"14511","Erzgebirgskreis":"14521","Mittelsachsen":"14522",
 "Vogtlandkreis":"14523","Zwickau":"14524","Dresden, Stadt":"14612",
 "Bautzen":"14625","Görlitz":"14626","Meißen":"14627",
 "Sächsische Schweiz-Osterzgebirge":"14628","Leipzig, Stadt":"14713",
 "Leipzig":"14729","Nordsachsen":"14730",
}
URL="https://www.statistik.sachsen.de/html/untergebrachte-wohnungslose-personen.html"
FILE="https://www.statistik.sachsen.de/download/soziales/statistik-sachsen_zr_untergebrachte_wohnungslose.xlsx"

def value(v):
    if not isinstance(v,(float,int)) or v<0 or int(v)!=v or int(v)%5:
        raise ValueError(f"Unexpected protected official count {v!r}")
    return int(v)

def main():
    hist=load_workbook(RAW/"saxony_homeless_2022_2026_district.xlsx",read_only=True,data_only=True).active
    current=load_workbook(RAW/"saxony_homeless_2026_district.xlsx",read_only=True,data_only=True).active
    h={}
    for row in hist.iter_rows(min_row=6,values_only=True):
        name,category = row[:2]
        if name in NAMES and str(category).strip()=="Insgesamt":
            if name in h:raise ValueError(f"Duplicate county {name}")
            h[name]={str(y):value(row[y-2022+2]) for y in range(2022,2027)}
    latest={}
    for row in current.iter_rows(min_row=6,max_row=18,values_only=True):
        if row[0] in NAMES:latest[row[0]]=value(row[1])
    assert len(h)==len(latest)==len(NAMES)==13,(len(h),len(latest))
    result=[]
    for name,ags in NAMES.items():
        if h[name]["2026"]!=latest[name]:
            raise ValueError(f"Current and historical source mismatch at {name}: {h[name]['2026']} vs {latest[name]}")
        result.append({"id":ags,"name":name,
                       "housed_homeless_2026_count":latest[name],
                       "annual_counts":h[name]})
    total=sum(r["housed_homeless_2026_count"] for r in result)
    state_total=value(current["B19"].value)
    if abs(total-state_total)>13*2:
        raise ValueError(f"Official state total mismatch (allow confidentiality rounding) {total} vs {state_total}")
    atlas=json.loads((WEB/"atlas-counties.json").read_text(encoding="utf-8"))
    keys={r["id"] for r in atlas["counties"]}
    assert all(r["id"] in keys for r in result)
    for d in (OUT,QA,WEB):d.mkdir(parents=True,exist_ok=True)
    with (OUT/"saxony_official_homeless_2022_2026.csv").open("w",encoding="utf-8",newline="") as f:
        writer=csv.writer(f)
        writer.writerow(["ags","county",*[str(y) for y in range(2022,2027)]])
        for r in sorted(result,key=lambda x:x["id"]):
            writer.writerow([r["id"],r["name"],*[r["annual_counts"][str(y)] for y in range(2022,2027)]])
    payload={"meta":{"publisher":"Statistisches Landesamt des Freistaates Sachsen","source":URL,
         "source_spreadsheet":FILE,"years":[2022,2023,2024,2025,2026],
         "reference_date":"31 January of each year","scope":"Only Saxony's 13 counties; other 15 states NOT sampled",
         "subject":"persons housed in official accommodation and qualifying placements; not rough sleepers or all homelessness",
         "rounding":"All counts rounded to multiples of 5 for statistical confidentiality; non-additivity is expected",
         "state_total_2026":state_total,"county_count":len(result)},
         "counties":sorted(result,key=lambda x:x["id"])}
    (WEB/"saxony-homeless-counties.json").write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":"))+"\n",encoding="utf-8")
    (QA/"saxony_homeless_audit.json").write_text(json.dumps({
       "district_count":len(result),"year_count":5,"min_year":2022,"max_year":2026,
       "sum_2026_county_rounded":total,"official_saxony_total_2026":state_total,
       "current_table_crosscheck_counties":len(latest),
       "crosswalk":NAMES,"source":URL,
       "warning":"Official rounded counts are disclosure-protected and cannot be treated as complete German homelessness."
    },ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("SUCCESS Saxony official housed homeless",len(result),"counties, 5 years, total",total)
if __name__=="__main__":main()
