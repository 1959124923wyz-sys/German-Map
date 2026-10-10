#!/usr/bin/env python3
"""Archive Saxony official 2026 district-level housed-homelessness tables.

The official statistical results are confidentiality-protected via rounding
to nearest 5, so district totals are not exactly additive.
Acquire the raw XLSX and make a workbook inventory. Parsing the district data
should happen only after inspecting this inventory (never guess cell position).
"""
from __future__ import annotations
import hashlib,json,urllib.request
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
FILES={
    "saxony_homeless_2026_district.xlsx":"https://www.statistik.sachsen.de/download/soziales/statistik-sachsen_untergebrachte_wohnungslose_kreise.xlsx",
    "saxony_homeless_2022_2026_district.xlsx":"https://www.statistik.sachsen.de/download/soziales/statistik-sachsen_zr_untergebrachte_wohnungslose.xlsx",
}
RAW=ROOT/"research/housing/raw"
QA=ROOT/"research/housing/qa"
PAGE="https://www.statistik.sachsen.de/html/untergebrachte-wohnungslose-personen.html"
def main():
    RAW.mkdir(parents=True,exist_ok=True);QA.mkdir(parents=True,exist_ok=True)
    output={"publisher":"Statistisches Landesamt des Freistaates Sachsen","source_landing":PAGE,
            "reference_day":"31 January 2026 (current) and series since 2022",
            "warning":"Official 5-person confidentiality rounding; counts do not add exactly; temporary accommodation only, not all persons experiencing homelessness.",
            "files":{}}
    for name,url in FILES.items():
        req=urllib.request.Request(url,headers={"User-Agent":"GermanMap housing research / official data collection"})
        with urllib.request.urlopen(req,timeout=80) as f:raw=f.read()
        if len(raw)<9000 or not raw.startswith(b"PK"):raise ValueError(f"Not a valid XLSX {name} size={len(raw)}")
        (RAW/name).write_bytes(raw)
        book=load_workbook(RAW/name,read_only=True,data_only=True)
        sheets={}
        for sheet in book:
            top=[]
            for row in sheet.iter_rows(min_row=1,max_row=min(sheet.max_row,28),max_col=min(sheet.max_column,20),values_only=True):
                top.append([str(v)[:180] if v is not None else None for v in row])
            sheets[sheet.title]={"rows":sheet.max_row,"cols":sheet.max_column,"first_28_rows":top}
            print("SHEET",name,sheet.title,sheet.max_row,sheet.max_column)
        book.close()
        output["files"][name]={"download":url,"sha256":hashlib.sha256(raw).hexdigest(),
             "bytes":len(raw),"sheets":sheets}
    (QA/"saxony_homeless_workbook_layout.json").write_text(json.dumps(output,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("SAXONY RAW SOURCES ARCHIVED; await schema review before publishing county values.")
if __name__=="__main__":main()
